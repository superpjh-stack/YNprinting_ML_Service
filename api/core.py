"""Local, evidence-preserving vision ML workbench."""
from pathlib import Path
import hashlib, json, os, sqlite3, shutil, unicodedata, uuid
from datetime import datetime, timezone
import numpy as np
from PIL import Image, ImageOps
from scipy import ndimage
from .collection import exclusions

ROOT = Path(os.environ.get('YN_RUNTIME', '.runtime')).resolve()
FEATURE_NAMES = ['brightness', 'contrast', 'edge_density', 'laplacian_variance', 'dark_fraction', 'light_fraction'] + [f'gray_hist_{i}' for i in range(16)]

def now(): return datetime.now(timezone.utc).isoformat()
def db():
    ROOT.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(ROOT / 'workbench.db'); c.row_factory = sqlite3.Row
    c.executescript('''CREATE TABLE IF NOT EXISTS samples(id TEXT PRIMARY KEY, data TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY, at TEXT, kind TEXT, data TEXT);
    CREATE TABLE IF NOT EXISTS runs(id TEXT PRIMARY KEY, data TEXT NOT NULL);''')
    return c

def all_samples():
    with db() as c: return [json.loads(r['data']) for r in c.execute('SELECT data FROM samples ORDER BY id')]
def save(s):
    with db() as c: c.execute('INSERT OR REPLACE INTO samples VALUES (?,?)', (s['id'], json.dumps(s, ensure_ascii=False)))
def event(kind, data):
    with db() as c: c.execute('INSERT INTO events(at,kind,data) VALUES (?,?,?)',(now(),kind,json.dumps(data,ensure_ascii=False)))
def sample(sid):
    with db() as c: r=c.execute('SELECT data FROM samples WHERE id=?',(sid,)).fetchone()
    if not r: raise ValueError('이미지를 찾을 수 없습니다.')
    return json.loads(r['data'])

def ingest(path, service='unassigned', group=None, weak_label='', references=None):
    path=Path(path)
    if path.stat().st_size > 30*1024*1024: raise ValueError('이미지 제한: 30MB')
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    with Image.open(path) as im:
        if im.width*im.height>40_000_000: raise ValueError('이미지 제한: 4천만 화소')
        im.verify()
    try: return sample(digest), False
    except ValueError: pass
    dest=ROOT/'raw'/f'{digest}{path.suffix.lower()}';dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(path,dest)
    s=dict(id=digest,name=path.name,path=str(dest),source=str(path),service=service,group=group or digest,
           weak_label=weak_label,references=references or [],label=None,reviewer=None,reviewed_at=None,
           domain='미확인',quality=None,features=None,created_at=now())
    with Image.open(path) as im:
        oriented=ImageOps.exif_transpose(im);s['image_width'],s['image_height']=oriented.size
    save(s);event('import',{'id':digest,'source':str(path)});return s,True

def import_folder(root):
    from openpyxl import load_workbook
    root=Path(root); mapping={}; warnings=[]
    for f in root.rglob('*.xlsx'):
        if '목록' not in unicodedata.normalize('NFC',f.name): continue
        wb=load_workbook(f,read_only=True,data_only=True)
        for row in wb.active.iter_rows(min_row=2,values_only=True):
            if len(row)>6 and row[1]: mapping[unicodedata.normalize('NFC',str(row[1]))]=str(row[6] or '')
        wb.close()
    count=0;duplicates=0
    for p in sorted(root.rglob('*')):
        if p.suffix.lower() not in ['.png','.jpg','.jpeg']:continue
        refs=[]
        for ref in sorted(p.parent.glob('*.pdf')):
            h=hashlib.sha256(ref.read_bytes()).hexdigest(); dest=ROOT/'references'/f'{h}.pdf';dest.parent.mkdir(parents=True,exist_ok=True)
            if not dest.exists():shutil.copy2(ref,dest)
            refs.append({'id':h,'name':ref.name,'path':str(dest),'status':'미검수'})
        try:
            _,created=ingest(p,group=unicodedata.normalize('NFC',p.parent.name),weak_label=mapping.get(unicodedata.normalize('NFC',p.parent.name),''),references=refs)
            count+=int(created);duplicates+=int(not created)
        except Exception as e: warnings.append({'file':str(p),'error':str(e)})
    result={'imported':count,'duplicates':duplicates,'errors':warnings};event('folder_import',result);return result

def extract(path):
    with Image.open(path) as im: rgb=ImageOps.exif_transpose(im).convert('RGB'); w,h=rgb.size;rgb.thumbnail((1024,1024));arr=np.asarray(rgb)
    gray=np.asarray(Image.fromarray(arr).convert('L'));g=gray/255.;edges=np.hypot(ndimage.sobel(g,axis=0),ndimage.sobel(g,axis=1))>0.3
    hist=np.histogram(gray,bins=16,range=(0,256))[0]/gray.size
    f=[g.mean(),g.std(),(edges>0).mean(),ndimage.laplace(gray.astype(float)).var(),(gray<15).mean(),(gray>240).mean(),*hist]
    return [float(x) for x in f],w,h

def process(stage):
    done=0;failed=0
    for s in all_samples():
        try:
            f,w,h=extract(s['path'])
            if stage=='clean':
                flags=[]
                if min(w,h)<128:flags.append('낮은 해상도')
                if f[3]<30:flags.append('흐림 검토')
                if f[0]<0.05 or f[0]>0.98:flags.append('노출 검토')
                s['quality']={'status':'review' if flags else 'pass','flags':flags,'width':w,'height':h,'version':'quality-v1'}
            elif stage=='features':
                if not s['quality']: continue
                s['features']={'version':'gray-edge-v1','values':f,'names':FEATURE_NAMES}
            save(s);done+=1
        except Exception as e:
            s['quality']={'status':'quarantine','flags':[str(e)]};s['features']=None;save(s);failed+=1
    result={'processed':done,'failed':failed};event(stage,result);return result

def readiness(service):
    rows=[s for s in all_samples() if s['service']==service]
    valid=[s for s in rows if s['label'] in ['OK','NG'] and s['reviewer'] and s['features'] and s['quality'] and s['quality']['status']=='pass' and not exclusions(s)]
    reasons=[]
    excluded=sum(bool(exclusions(s)) for s in rows)
    if excluded:reasons.append(f'{excluded}장은 결함 표시 또는 촬영 대상 검수가 필요합니다.')
    if len(valid)<12:reasons.append('검수·정제·피처 완료 표본이 최소 12개 필요합니다 (개발용 최소 조건).')
    if len(set(s['group'] for s in valid))<6:reasons.append('독립 그룹이 최소 6개 필요합니다.')
    for label in ['OK','NG']:
        if len({s['group'] for s in valid if s['label']==label})<3:reasons.append(f'{label} 표본이 최소 3개 독립 그룹에 필요합니다.')
    return {'service':service,'total':len(rows),'eligible':len(valid),'reasons':reasons,'ready':not reasons},valid

def overview():
    rows=all_samples()
    with db() as c:runs=[json.loads(x['data']) for x in c.execute('SELECT data FROM runs ORDER BY rowid DESC')]
    return {'total':len(rows),'unassigned':sum(s['service']=='unassigned' for s in rows),
      'cleaned':sum(bool(s['quality']) for s in rows),'featured':sum(bool(s['features']) for s in rows),
      'overlay_pending':sum(s.get('overlay_status','unknown')=='unknown' for s in rows),'reviewed':sum(bool(s['reviewer']) for s in rows),'services':[readiness(x)[0] for x in ['plate','die']], 'runs':runs,
      'feature_names':FEATURE_NAMES,'limitation':'전체 이미지 OK/NG 기준모델. OCR·결함 위치 검출은 아직 연결되지 않았습니다.'}
